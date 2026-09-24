#!/usr/bin/env python3
"""
Scan a directory for image files (.png, .jpg, .jpeg) and create a PowerPoint
presentation with one image per slide, centered and scaled to fit the slide.

Usage:
    python3 images_to_pptx.py [directory] [-o output.pptx] [-r]

    directory   Directory to scan for images (default: current directory)
    -o, --output  Output .pptx file path (default: <directory_name>_images.pptx)
    -r, --recursive  Recurse into subdirectories
"""
import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image

IMAGE_EXTS = {".png", ".jpg", ".jpeg"}


def find_images(directory: Path, recursive: bool):
    pattern = "**/*" if recursive else "*"
    files = [
        p for p in directory.glob(pattern)
        if p.is_file() and p.suffix.lower() in IMAGE_EXTS
    ]
    return sorted(files)


def add_image_slide(prs: Presentation, image_path: Path, input_dir: Path, output_file: Path):
    blank_slide_layout = prs.slide_layouts[6]  # blank layout
    slide = prs.slides.add_slide(blank_slide_layout)

    slide_w, slide_h = prs.slide_width, prs.slide_height

    with Image.open(image_path) as img:
        img_w, img_h = img.size

    # Scale image to fit within slide while preserving aspect ratio
    scale = min(slide_w / img_w, slide_h / img_h)
    new_w = int(img_w * scale)
    new_h = int(img_h * scale)

    left = (slide_w - new_w) // 2
    top = (slide_h - new_h) // 2

    slide.shapes.add_picture(str(image_path), left, top, width=new_w, height=new_h)
    
    # Add text note with input and output file information at the bottom
    text_box = slide.shapes.add_textbox(
        Inches(0.25),
        slide_h - Inches(0.75),
        slide_w - Inches(0.5),
        Inches(0.5)
    )
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    note_text = f"Input: {input_dir.name} | Output: {output_file.name} | Image: {image_path.name}"
    p = text_frame.paragraphs[0]
    p.text = note_text
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(128, 128, 128)  # Gray color
    p.alignment = PP_ALIGN.LEFT


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", nargs="?", default=".", help="Directory to scan for images")
    parser.add_argument("-o", "--output", help="Output .pptx file path")
    parser.add_argument("-r", "--recursive", action="store_true", help="Recurse into subdirectories")
    args = parser.parse_args()

    directory = Path(args.directory).expanduser().resolve()
    if not directory.is_dir():
        print(f"Error: '{directory}' is not a directory", file=sys.stderr)
        sys.exit(1)

    images = find_images(directory, args.recursive)
    if not images:
        print(f"No .png/.jpg/.jpeg images found in '{directory}'")
        sys.exit(0)

    output = Path(args.output).expanduser().resolve() if args.output else directory / f"{directory.name}_images.pptx"

    prs = Presentation()
    # Use a widescreen 16:9 slide size
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for img_path in images:
        try:
            add_image_slide(prs, img_path, directory, output)
            print(f"Added: {img_path.name}")
        except Exception as e:
            print(f"Skipping '{img_path.name}': {e}", file=sys.stderr)

    prs.save(output)
    print(f"\nSaved {len(images)} image(s) to: {output}")


if __name__ == "__main__":
    main()
